from io import BytesIO
from docx import Document
from fpdf import FPDF
def format_txt(text):
    """Return document text as UTF-8 bytes."""
    return str(text).encode("utf-8")
def format_docx(text, draft_type="Document"):
    """Create a DOCX file in memory."""
    document = Document()
    document.add_heading(str(draft_type), level=1)
    for paragraph in str(text).splitlines():
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
def format_pdf(text, draft_type="Document"):
    """Create a PDF file in memory."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 10, str(draft_type))
    pdf.ln(3)
    for paragraph in str(text).splitlines():
        safe_paragraph = paragraph.encode(
            "latin-1", errors="replace"
        ).decode("latin-1")
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 8, safe_paragraph)
    return bytes(pdf.output())
